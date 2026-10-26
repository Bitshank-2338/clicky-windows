import asyncio
import io
import json
import wave
import zipfile
from unittest.mock import Mock

import pytest


def test_offline_voice_generates_real_wav(monkeypatch):
    import audio.tts.sapi_tts_provider as sapi
    captured = []
    async def play(content):
        captured.append(content)
    monkeypatch.setattr(sapi,'play_mp3_async',play)
    asyncio.run(sapi.SapiTTSProvider().speak('Clicky offline speech test.'))
    with wave.open(io.BytesIO(captured[0])) as audio:
        assert audio.getnframes() > 1000 and audio.getnchannels() == 1


def test_xlsx_and_word_tables(tmp_path):
    from tutor_features.pdf_context import extract_text
    from docx import Document
    word = tmp_path / 'table.docx'
    doc = Document()
    doc.add_table(1,1).cell(0,0).text = 'Table content must survive'
    doc.save(word)
    assert 'Table content must survive' in extract_text(word)
    sheet = tmp_path / 'sample.xlsx'
    with zipfile.ZipFile(sheet,'w') as archive:
        archive.writestr('xl/worksheets/sheet1.xml','<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row><c r="A1" t="inlineStr"><is><t>Hello</t></is></c><c r="B1"><v>42</v></c></row></sheetData></worksheet>')
    text = extract_text(sheet)
    assert 'A1: Hello' in text and 'B1: 42' in text


def test_conversations_boundaries_and_clear():
    import session_store
    from ai.base_provider import Message
    session_store.save('editor',[Message('user','explain this'),Message('assistant','answer')])
    assert session_store.load('editor')[0]['content'] == 'explain this'
    assert session_store.load('browser') == []
    session_store.clear()
    assert session_store.load('editor') == []


def test_speech_starts_before_generation_finishes_and_stop_cancels(monkeypatch,qt_app):
    import threading
    import companion_manager as module
    monkeypatch.setattr(module.skills_pkg,'load_all',lambda:[])
    monkeypatch.setattr(module,'active_window_title',lambda:'Test editor')
    started,finished = threading.Event(),threading.Event()
    class LLM:
        async def stream_response(self,**kwargs):
            try:
                yield 'First sentence. Second sentence'
                await asyncio.sleep(30)
                yield ' finished.'
            finally:finished.set()
    class TTS:
        async def speak(self,text):started.set()
    manager=module.CompanionManager()
    manager._journal_enabled=False
    manager._get_llm=lambda:LLM()
    manager._get_tts=lambda:TTS()
    try:
        manager.submit_text('who is Ada Lovelace')
        assert started.wait(timeout=5), 'Narration waited for the whole answer'
        assert not finished.is_set()
        manager.stop()
        assert finished.wait(timeout=5)
    finally:
        manager.shutdown();manager._thread.join(timeout=5)


def test_job_claim_recovery_and_routine(monkeypatch):
    from jobs import JobStore
    store = JobStore()
    ident = store.add('Report','Write a short report',interval=60,next_run=0)
    assert store.claim(now=1)['id'] == ident
    assert store.claim(now=1) is None
    store.finish(ident,result='A useful result')
    row = store.list()[0]
    assert row['state']=='completed' and row['next_run']
    assert store.claim(now=row['next_run']+1)['id']==ident
    store.recover()
    assert store.list()[0]['state']=='interrupted'
    store.action(ident,'retry')
    assert store.list()[0]['state']=='queued'


def test_recording_muxes_narration_without_capturing_desktop(monkeypatch):
    import time
    import numpy as np
    import av
    import tutor_features.lesson_recorder as module
    import screen.privacy
    class FakeCapture:
        monitors=[{}, {'left':0,'top':0,'width':64,'height':48}]
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def grab(self,monitor):
            from types import SimpleNamespace
            return SimpleNamespace(size=(64,48),bgra=bytes([40,80,160,255])*64*48)
    monkeypatch.setattr(module.mss,'mss',FakeCapture)
    monkeypatch.setattr(screen.privacy,'redact',lambda image,*args:image)
    recorder=module.LessonRecorder()
    path=recorder.start()
    time.sleep(.2)
    recorder._record_audio(np.ones(2400,dtype=np.float32)*.1,24000)
    time.sleep(.15)
    recorder.stop()
    with av.open(str(path / 'lesson-narrated.mp4')) as video:
        assert video.streams.video and video.streams.audio
        assert video.streams.video[0].width==64


def test_mcp_http_session_and_authorization(monkeypatch):
    from aiohttp import web
    from connectors import MCPConnection, save_connector
    received = []
    async def scenario():
        async def handler(request):
            payload = await request.json()
            received.append((payload,dict(request.headers)))
            method = payload['method']
            if method=='notifications/initialized':return web.Response(status=202)
            if method=='initialize':result={'protocolVersion':'2025-06-18','capabilities':{},'serverInfo':{'name':'test','version':'1'}}
            elif method=='tools/list':result={'tools':[{'name':'echo','inputSchema':{'type':'object'}}]}
            else:result={'content':[{'type':'text','text':payload['params']['arguments']['text']}]}
            return web.json_response({'jsonrpc':'2.0','id':payload['id'],'result':result},headers={'Mcp-Session-Id':'test-session'})
        app=web.Application();app.router.add_post('/mcp',handler)
        runner=web.AppRunner(app);await runner.setup()
        site=web.TCPSite(runner,'127.0.0.1',0);await site.start()
        port=site._server.sockets[0].getsockname()[1]
        try:
            ident=save_connector('test',f'http://127.0.0.1:{port}/mcp','test-token')
            connection=MCPConnection({'id':ident,'url':f'http://127.0.0.1:{port}/mcp','name':'test'})
            assert (await connection.call('echo',{'text':'ok'}))['content'][0]['text']=='ok'
            with pytest.raises(ValueError):await connection.call('not-listed',{})
        finally:await runner.cleanup()
    asyncio.run(scenario())
    assert received[0][1]['Authorization']=='Bearer test-token'
    assert received[-1][1]['Mcp-Session-Id']=='test-session'


def test_cloud_transcription_respects_selected_language(monkeypatch):
    import httpx
    import audio.stt.deepgram_stt as module
    captured=[]
    class Client:
        def __init__(self,**kwargs):pass
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        async def post(self,url,**kwargs):
            captured.append(kwargs['params'])
            return httpx.Response(200,json={'results':{'channels':[{'alternatives':[{'transcript':'नमस्ते'}]}]}},request=httpx.Request('POST',url))
    monkeypatch.setattr(module.httpx,'AsyncClient',Client)
    monkeypatch.setattr(module.cfg,'whisper_language','hi')
    monkeypatch.setattr(module.cfg,'deepgram_api_key','test-key')
    assert asyncio.run(module.DeepgramSTT().transcribe(b'\0\0'*1600))=='नमस्ते'
    assert captured[0]['language']=='hi' and captured[0]['model']=='nova-3'


def test_live_transcription_streams_pcm_and_partial_results(monkeypatch):
    import aiohttp
    import audio.stt.deepgram_stt as module
    from types import SimpleNamespace
    partials=[]
    class Socket:
        def __init__(self):self.queue=asyncio.Queue();self.count=0
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        def __aiter__(self):return self
        async def __anext__(self):
            item=await self.queue.get()
            if item is None:raise StopAsyncIteration
            return item
        async def send_bytes(self,content):
            assert content==b'\0\0'*10
            self.count+=1
            for text,final in [('partial',False),('Hello' if self.count==1 else 'world',True)]:
                await self.queue.put(SimpleNamespace(type=aiohttp.WSMsgType.TEXT,data=json.dumps({'type':'Results','is_final':final,'channel':{'alternatives':[{'transcript':text}]}})))
        async def send_json(self,value):
            assert value['type']=='CloseStream'
            await self.queue.put(None)
    class Session:
        def __init__(self,**kwargs):pass
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        def ws_connect(self,url,headers):
            assert 'nova-3' in url and 'Authorization' in headers
            return Socket()
    async def chunks():
        yield b'\0\0'*10
        yield b'\0\0'*10
    monkeypatch.setattr(module.aiohttp,'ClientSession',Session)
    monkeypatch.setattr(module.cfg,'deepgram_api_key','test-key')
    assert asyncio.run(module.DeepgramSTT().transcribe_stream(chunks(),partials.append))=='Hello world'
    assert partials[0]=='partial' and partials[-1]=='Hello world'


def test_main_wires_all_windows_without_network_or_hooks(qt_app,monkeypatch):
    import main
    import keyboard
    import desktop
    import companion_manager
    monkeypatch.setattr(main,'_setup_logging',lambda:None)
    monkeypatch.setattr(companion_manager.skills_pkg,'load_all',lambda:[])
    monkeypatch.setattr(main.GlobalHotkeyMonitor,'start',lambda self:None)
    monkeypatch.setattr(main.StopHotkey,'start',lambda self:None)
    monkeypatch.setattr(keyboard,'add_hotkey',lambda *args,**kwargs:None)
    monkeypatch.setattr(desktop,'foreground',lambda:{'hwnd':123,'pid':0,'title':'Test editor','point':(0,0)})
    monkeypatch.setattr(main.CompanionManager,'start',lambda self:None)
    monkeypatch.setattr(main.CompanionManager,'refresh_ollama_models',lambda self:None)
    import ui.setup_wizard
    monkeypatch.setattr(ui.setup_wizard,'maybe_show_setup_wizard',lambda:None)
    monkeypatch.setattr(qt_app,'exec',lambda:0)
    created=[]
    original=main.CompanionManager
    def manager():
        instance=original();created.append(instance);return instance
    monkeypatch.setattr(main,'CompanionManager',manager)
    try:
        with pytest.raises(SystemExit) as exit_result:main.main()
        assert exit_result.value.code==0
        titles={window.windowTitle() for window in qt_app.topLevelWidgets()}
        assert 'Clicky Settings' in titles and 'Clicky Connections' in titles
    finally:
        for instance in created:instance.shutdown();instance._thread.join(timeout=5)
        for widget in qt_app.topLevelWidgets():widget.close()
